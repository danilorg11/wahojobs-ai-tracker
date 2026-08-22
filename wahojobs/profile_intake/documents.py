"""In-memory, deterministic PDF and DOCX resume text extraction."""

from __future__ import annotations

from io import BytesIO
import re
import unicodedata
from xml.etree import ElementTree
import zipfile

from wahojobs.profile_intake.contracts import (
    DEFAULT_DOCUMENT_LIMITS,
    DocumentFormat,
    DocumentKind,
    DocumentLimits,
    EvidenceBlock,
    ExtractedDocument,
    ParserMetadata,
    ProfileIntakeError,
    _require_document_reference,
)


_PDF_HEADER = b"%PDF-"
_LEGACY_OFFICE_HEADER = bytes.fromhex("d0cf11e0a1b11ae1")
_DOCX_MAIN_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
)
_SUPPORTED_ZIP_COMPRESSION = frozenset({zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED})


def normalize_extracted_text(text: object) -> str:
    """Conservatively normalize parser text without interpreting its content."""

    if type(text) is not str:
        raise ProfileIntakeError("invalid_extracted_text")
    text = unicodedata.normalize("NFC", text.replace("\r\n", "\n").replace("\r", "\n"))
    cleaned: list[str] = []
    for character in text:
        if character in ("\n", "\t"):
            cleaned.append(character)
        elif character == "\u00a0":
            cleaned.append(" ")
        elif unicodedata.category(character) != "Cc":
            cleaned.append(character)
    lines = [re.sub(r"[^\S\n]+", " ", line).strip() for line in "".join(cleaned).split("\n")]
    normalized_lines: list[str] = []
    previous_blank = True
    for line in lines:
        if line:
            normalized_lines.append(line)
            previous_blank = False
        elif not previous_blank:
            normalized_lines.append("")
            previous_blank = True
    while normalized_lines and not normalized_lines[-1]:
        normalized_lines.pop()
    return "\n".join(normalized_lines)


def _split_oversized_segment(segment: str, max_chars: int) -> list[str]:
    pieces: list[str] = []
    remaining = segment
    while len(remaining) > max_chars:
        split_at = remaining.rfind(" ", 0, max_chars + 1)
        if split_at < max_chars // 2:
            split_at = max_chars
        piece = remaining[:split_at].strip()
        if piece:
            pieces.append(piece)
        remaining = remaining[split_at:].strip()
    if remaining:
        pieces.append(remaining)
    return pieces


def build_evidence_blocks(
    normalized_text: object,
    *,
    limits: DocumentLimits = DEFAULT_DOCUMENT_LIMITS,
) -> tuple[EvidenceBlock, ...]:
    """Pack normalized paragraphs into stable, bounded numbered blocks."""

    if type(limits) is not DocumentLimits:
        raise ProfileIntakeError("invalid_document_limits")
    if type(normalized_text) is not str or not normalized_text:
        raise ProfileIntakeError("no_extractable_text")
    if len(normalized_text) > limits.max_normalized_text_chars:
        raise ProfileIntakeError("extracted_text_too_large")
    segments: list[str] = []
    for paragraph in re.split(r"\n\s*\n", normalized_text):
        paragraph = paragraph.strip()
        if paragraph:
            segments.extend(_split_oversized_segment(paragraph, limits.max_evidence_block_chars))

    packed: list[str] = []
    current = ""
    for segment in segments:
        candidate = segment if not current else f"{current}\n\n{segment}"
        if len(candidate) <= limits.max_evidence_block_chars:
            current = candidate
            continue
        if current:
            packed.append(current)
        current = segment
    if current:
        packed.append(current)
    if not packed:
        raise ProfileIntakeError("no_extractable_text")
    if len(packed) > limits.max_evidence_blocks:
        raise ProfileIntakeError("evidence_block_limit_exceeded")
    return tuple(
        EvidenceBlock(reference=f"b{index:03d}", text=text)
        for index, text in enumerate(packed, start=1)
    )


def _validate_input(
    document_bytes: object,
    document_reference: object,
    limits: object,
) -> bytes:
    _require_document_reference(document_reference)
    if type(limits) is not DocumentLimits:
        raise ProfileIntakeError("invalid_document_limits")
    if type(document_bytes) is not bytes or not document_bytes:
        raise ProfileIntakeError("invalid_document_bytes")
    if len(document_bytes) > limits.max_upload_bytes:
        raise ProfileIntakeError(
            "upload_too_large",
            diagnostics={"byte_count": len(document_bytes)},
        )
    return document_bytes


def _finalize_document(
    *,
    document_bytes: bytes,
    document_reference: str,
    document_kind: DocumentKind,
    document_format: DocumentFormat,
    page_count: int | None,
    parser_name: str,
    parser_version: str,
    text: str,
    limits: DocumentLimits,
) -> ExtractedDocument:
    normalized = normalize_extracted_text(text)
    if len(normalized) > limits.max_normalized_text_chars:
        raise ProfileIntakeError(
            "extracted_text_too_large",
            diagnostics={"character_count": len(normalized)},
        )
    alphanumeric_count = sum(character.isalnum() for character in normalized)
    if alphanumeric_count < limits.min_extractable_alphanumeric_chars:
        raise ProfileIntakeError("no_extractable_text")
    blocks = build_evidence_blocks(normalized, limits=limits)
    return ExtractedDocument(
        document_reference=document_reference,
        document_kind=document_kind,
        document_format=document_format,
        original_byte_size=len(document_bytes),
        normalized_text_chars=len(normalized),
        page_count=page_count,
        parser=ParserMetadata(parser=parser_name, version=parser_version),
        evidence_blocks=blocks,
    )


def _extract_pdf(
    document_bytes: bytes,
    *,
    document_reference: str,
    document_kind: DocumentKind,
    limits: DocumentLimits,
) -> ExtractedDocument:
    if _PDF_HEADER not in document_bytes[:1024]:
        raise ProfileIntakeError("invalid_pdf")
    try:
        import pypdf
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(document_bytes), strict=True)
        if reader.is_encrypted:
            raise ProfileIntakeError("encrypted_pdf")
        page_count = len(reader.pages)
        if page_count < 1:
            raise ProfileIntakeError("no_extractable_text")
        if page_count > limits.max_pdf_pages:
            raise ProfileIntakeError(
                "pdf_page_limit_exceeded",
                diagnostics={"page_count": page_count},
            )
        page_text: list[str] = []
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                page_text.append(extracted)
        return _finalize_document(
            document_bytes=document_bytes,
            document_reference=document_reference,
            document_kind=document_kind,
            document_format=DocumentFormat.PDF,
            page_count=page_count,
            parser_name="pypdf",
            parser_version=pypdf.__version__,
            text="\n\n".join(page_text),
            limits=limits,
        )
    except ProfileIntakeError:
        raise
    except Exception:
        raise ProfileIntakeError("invalid_pdf") from None


def _is_traversal_member(name: str) -> bool:
    normalized = name.replace("\\", "/")
    return (
        not normalized
        or normalized.startswith("/")
        or re.match(r"^[A-Za-z]:", normalized) is not None
        or any(part == ".." for part in normalized.split("/"))
    )


def _read_bounded_member(
    archive: zipfile.ZipFile,
    info: zipfile.ZipInfo,
    *,
    max_bytes: int,
    size_error_code: str = "docx_relationships_too_large",
) -> bytes:
    if info.file_size > max_bytes:
        raise ProfileIntakeError(size_error_code)
    try:
        content = archive.read(info)
    except Exception:
        raise ProfileIntakeError("invalid_docx") from None
    if len(content) > max_bytes:
        raise ProfileIntakeError(size_error_code)
    return content


def _contains_unsafe_xml_declaration(content: bytes) -> bool:
    compact = content.lower().replace(b"\x00", b"")
    return b"<!doctype" in compact or b"<!entity" in compact


def _preflight_docx(document_bytes: bytes, limits: DocumentLimits) -> None:
    if document_bytes.startswith(_LEGACY_OFFICE_HEADER):
        raise ProfileIntakeError("unsupported_office_format")
    if not document_bytes.startswith(b"PK") or not zipfile.is_zipfile(BytesIO(document_bytes)):
        raise ProfileIntakeError("invalid_docx")
    try:
        with zipfile.ZipFile(BytesIO(document_bytes)) as archive:
            members = archive.infolist()
            if len(members) > limits.max_docx_members:
                raise ProfileIntakeError(
                    "docx_member_limit_exceeded",
                    diagnostics={"member_count": len(members)},
                )
            names = [info.filename for info in members]
            if len(names) != len(set(names)):
                raise ProfileIntakeError("invalid_docx")
            total_uncompressed = 0
            total_compressed = 0
            by_name: dict[str, zipfile.ZipInfo] = {}
            for info in members:
                if _is_traversal_member(info.filename):
                    raise ProfileIntakeError("docx_path_traversal")
                if info.flag_bits & 0x1:
                    raise ProfileIntakeError("encrypted_docx_member")
                if info.compress_type not in _SUPPORTED_ZIP_COMPRESSION:
                    raise ProfileIntakeError("unsupported_docx_compression")
                if info.file_size > limits.max_docx_member_bytes:
                    raise ProfileIntakeError("docx_member_too_large")
                total_uncompressed += info.file_size
                total_compressed += info.compress_size
                if total_uncompressed > limits.max_docx_uncompressed_bytes:
                    raise ProfileIntakeError(
                        "docx_uncompressed_size_exceeded",
                        diagnostics={"uncompressed_bytes": total_uncompressed},
                    )
                if (
                    info.file_size > 0
                    and info.file_size / max(1, info.compress_size)
                    > limits.max_docx_compression_ratio
                ):
                    raise ProfileIntakeError("docx_compression_ratio_exceeded")
                by_name[info.filename] = info
            if (
                total_uncompressed > 0
                and total_uncompressed / max(1, total_compressed)
                > limits.max_docx_compression_ratio
            ):
                raise ProfileIntakeError("docx_compression_ratio_exceeded")
            if "[Content_Types].xml" not in by_name or "word/document.xml" not in by_name:
                raise ProfileIntakeError("invalid_docx")
            lowered_names = {name.casefold() for name in names}
            if any(name.endswith("vbaproject.bin") for name in lowered_names):
                raise ProfileIntakeError("unsupported_macro_container")

            for info in members:
                lowered_name = info.filename.casefold()
                if not (lowered_name.endswith(".xml") or lowered_name.endswith(".rels")):
                    continue
                xml_content = _read_bounded_member(
                    archive,
                    info,
                    max_bytes=limits.max_docx_member_bytes,
                    size_error_code="docx_member_too_large",
                )
                if _contains_unsafe_xml_declaration(xml_content):
                    raise ProfileIntakeError("unsafe_docx_xml")

            content_types = _read_bounded_member(
                archive,
                by_name["[Content_Types].xml"],
                max_bytes=limits.max_relationship_xml_bytes,
            )
            try:
                content_root = ElementTree.fromstring(content_types)
            except ElementTree.ParseError:
                raise ProfileIntakeError("invalid_docx") from None
            main_content_type = None
            for element in content_root.iter():
                if element.tag.rsplit("}", 1)[-1] != "Override":
                    continue
                content_type = element.attrib.get("ContentType", "")
                if "macroenabled" in content_type.casefold() or "vbaproject" in content_type.casefold():
                    raise ProfileIntakeError("unsupported_macro_container")
                if element.attrib.get("PartName") == "/word/document.xml":
                    main_content_type = content_type
            if main_content_type != _DOCX_MAIN_CONTENT_TYPE:
                raise ProfileIntakeError("invalid_docx")

            for info in members:
                if not info.filename.casefold().endswith(".rels"):
                    continue
                relationship_xml = _read_bounded_member(
                    archive,
                    info,
                    max_bytes=limits.max_relationship_xml_bytes,
                )
                try:
                    relationship_root = ElementTree.fromstring(relationship_xml)
                except ElementTree.ParseError:
                    raise ProfileIntakeError("invalid_docx") from None
                for relationship in relationship_root.iter():
                    if relationship.attrib.get("TargetMode", "").casefold() == "external":
                        raise ProfileIntakeError("docx_external_relationship")
    except ProfileIntakeError:
        raise
    except (zipfile.BadZipFile, OSError, ValueError):
        raise ProfileIntakeError("invalid_docx") from None


def _table_text(table) -> list[str]:
    lines: list[str] = []
    seen_cells: set[object] = set()
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            cell_identity = cell._tc
            if cell_identity in seen_cells:
                continue
            seen_cells.add(cell_identity)
            cell_text = "\n".join(
                paragraph.text for paragraph in cell.paragraphs if paragraph.text.strip()
            ).strip()
            if cell_text:
                cells.append(cell_text)
            for nested_table in cell.tables:
                cells.extend(_table_text(nested_table))
        if cells:
            lines.append(" | ".join(cells))
    return lines


def _container_text(container) -> list[str]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    lines: list[str] = []
    try:
        blocks = container.iter_inner_content()
    except AttributeError:
        blocks = (*container.paragraphs, *container.tables)
    for block in blocks:
        if isinstance(block, Paragraph) and block.text.strip():
            lines.append(block.text)
        elif isinstance(block, Table):
            lines.extend(_table_text(block))
    return lines


def _extract_docx(
    document_bytes: bytes,
    *,
    document_reference: str,
    document_kind: DocumentKind,
    limits: DocumentLimits,
) -> ExtractedDocument:
    _preflight_docx(document_bytes, limits)
    try:
        import docx
        from docx import Document

        document = Document(BytesIO(document_bytes))
        header_text: list[str] = []
        footer_text: list[str] = []
        seen_headers: set[str] = set()
        seen_footers: set[str] = set()
        for section in document.sections:
            for header in (section.header, section.first_page_header, section.even_page_header):
                part_name = str(header.part.partname)
                if part_name not in seen_headers:
                    seen_headers.add(part_name)
                    header_text.extend(_container_text(header))
            for footer in (section.footer, section.first_page_footer, section.even_page_footer):
                part_name = str(footer.part.partname)
                if part_name not in seen_footers:
                    seen_footers.add(part_name)
                    footer_text.extend(_container_text(footer))
        body_text = _container_text(document)
        return _finalize_document(
            document_bytes=document_bytes,
            document_reference=document_reference,
            document_kind=document_kind,
            document_format=DocumentFormat.DOCX,
            page_count=None,
            parser_name="python-docx",
            parser_version=docx.__version__,
            text="\n\n".join((*header_text, *body_text, *footer_text)),
            limits=limits,
        )
    except ProfileIntakeError:
        raise
    except Exception:
        raise ProfileIntakeError("invalid_docx") from None


def extract_profile_document(
    document_bytes: object,
    *,
    document_reference: object,
    document_kind: object,
    document_format: object,
    limits: DocumentLimits = DEFAULT_DOCUMENT_LIMITS,
) -> ExtractedDocument:
    """Extract one bounded profile document without retaining its bytes."""

    content = _validate_input(document_bytes, document_reference, limits)
    if type(document_kind) is not DocumentKind:
        raise ProfileIntakeError("invalid_document_kind")
    if type(document_format) is not DocumentFormat:
        raise ProfileIntakeError("invalid_document_format")
    if (
        document_kind is DocumentKind.LINKEDIN_PROFILE_EXPORT
        and document_format is not DocumentFormat.PDF
    ):
        raise ProfileIntakeError("unsupported_document_origin_format")
    if document_format is DocumentFormat.PDF:
        return _extract_pdf(
            content,
            document_reference=document_reference,
            document_kind=document_kind,
            limits=limits,
        )
    if document_format is DocumentFormat.DOCX:
        return _extract_docx(
            content,
            document_reference=document_reference,
            document_kind=document_kind,
            limits=limits,
        )
    raise ProfileIntakeError("unsupported_document_format")


def extract_resume_document(
    document_bytes: object,
    *,
    document_reference: object,
    document_format: object,
    limits: DocumentLimits = DEFAULT_DOCUMENT_LIMITS,
) -> ExtractedDocument:
    """Backward-compatible resume-only extraction boundary."""

    return extract_profile_document(
        document_bytes,
        document_reference=document_reference,
        document_kind=DocumentKind.RESUME,
        document_format=document_format,
        limits=limits,
    )
